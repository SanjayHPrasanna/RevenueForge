"""Company profile, onboarding uploads and the knowledge library (capabilities, people, certifications,
rate card, case studies, accelerators, proposal template)."""
from __future__ import annotations

import io
import re
import secrets
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import delete, func, select

from ..config import UPLOAD_DIR, settings
from ..models import Accelerator, CaseStudy, Capability, Certification, Document, Employee, RateCard, Tenant
from ..security import Ctx, current_ctx, require_roles
from ..services import folder_import, graph, llm
from ..services.onboarding import FIELDS, KINDS, analyze, missing_fields, records_from_table
from ..services.parsing import parse_document

router = APIRouter(prefix="/api", tags=["library"])

MODEL = {"capabilities": Capability, "skills": Employee, "certifications": Certification, "rates": RateCard,
         "case_studies": CaseStudy, "accelerators": Accelerator}
ONBOARDING_STEPS = ["profile", "capabilities", "skills", "certifications", "rates", "case_studies", "template"]
ALLOWED = {".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".pptx", ".txt", ".md"}


def _row_dict(kind: str, r) -> dict:
    cols = [c.name for c in MODEL[kind].__table__.columns if c.name not in ("tenant_id",)]
    return {c: getattr(r, c) for c in cols}


async def save_upload(ctx: Ctx, file: UploadFile, kind: str) -> tuple[Document, bytes]:
    ext = Path(file.filename or "").suffix.lower()
    if ext not in ALLOWED:
        raise HTTPException(400, f"{ext or 'This file type'} isn't supported. Upload PDF, Word, Excel, CSV, PowerPoint or text.")
    data = await file.read()
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, f"File is larger than {settings.max_upload_mb} MB.")
    if not data:
        raise HTTPException(400, "The file is empty.")
    doc = Document(tenant_id=ctx.tenant_id, kind=kind, filename=file.filename, path="", size=len(data), uploaded_by=ctx.user.id)
    ctx.db.add(doc)
    ctx.db.flush()
    folder = UPLOAD_DIR / ctx.tenant_id
    folder.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9._ -]+", "_", file.filename)[-120:]
    path = folder / f"{doc.id}_{safe}"
    path.write_bytes(data)
    doc.path = str(path)
    return doc, data


def _onboarding_status(ctx: Ctx) -> dict:
    t = ctx.tenant
    counts = {k: ctx.db.scalar(select(func.count()).select_from(m).where(m.tenant_id == ctx.tenant_id)) for k, m in MODEL.items()}
    has_template = bool(ctx.db.scalar(select(func.count()).select_from(Document).where(
        Document.tenant_id == ctx.tenant_id, Document.kind == "template", Document.status == "imported")))
    steps = {
        "profile": bool(t.about and t.industry),
        "capabilities": counts["capabilities"] > 0, "skills": counts["skills"] > 0, "certifications": counts["certifications"] > 0,
        "rates": counts["rates"] > 0, "case_studies": counts["case_studies"] > 0, "template": has_template,
    }
    required = ["capabilities", "skills", "rates", "template"]
    return {"steps": steps, "counts": counts, "complete": all(steps[s] for s in required), "required": required,
            "pct": round(sum(1 for s in ONBOARDING_STEPS if steps[s]) / len(ONBOARDING_STEPS) * 100)}


@router.get("/company")
def get_company(ctx: Ctx = Depends(current_ctx)):
    t = ctx.tenant
    return {"id": t.id, "name": t.name, "slug": t.slug, "industry": t.industry, "website": t.website, "about": t.about,
            "headquarters": t.headquarters, "headcount": t.headcount, "founded": t.founded, "currency": t.currency,
            "settings": {k: v for k, v in (t.settings or {}).items() if k != "ai"}, "onboarding": _onboarding_status(ctx),
            "llm": llm.status(), "notice": llm.notice()}


# --------------------------------------------------------------------------- the company's own AI provider and key
@router.get("/company/ai")
def get_company_ai(ctx: Ctx = Depends(current_ctx)):
    from ..services import company_ai
    return company_ai.view(ctx.tenant)


@router.put("/company/ai")
def save_company_ai(body: dict, ctx: Ctx = Depends(current_ctx)):
    """Anyone in the company may change it; it applies to this company only."""
    from ..services import company_ai
    try:
        company_ai.save(ctx.db, ctx.tenant, body, ctx.user.email)
    except ValueError as e:
        raise HTTPException(400, str(e))
    ctx.audit("company.ai.update", str(body.get("provider", "")))
    ctx.db.commit()
    return company_ai.view(ctx.tenant)


@router.post("/company/ai/test")
def test_company_ai(ctx: Ctx = Depends(current_ctx)):
    from ..services import company_ai
    return company_ai.test(ctx.tenant_id)


@router.put("/company")
def update_company(body: dict, ctx: Ctx = Depends(require_roles())):
    t = ctx.tenant
    for f in ("name", "industry", "website", "about", "headquarters", "founded", "currency"):
        if f in body and body[f] is not None:
            setattr(t, f, str(body[f])[:2000])
    if "headcount" in body:
        try:
            t.headcount = int(body["headcount"] or 0)
        except ValueError:
            pass
    if isinstance(body.get("settings"), dict):
        s = dict(t.settings or {})
        for k, v in body["settings"].items():
            if k == "ai":
                continue  # the company's AI provider and key are changed only through /company/ai
            if k in ("target_margin", "min_margin", "travel_pct"):
                v = float(v)
            s[k] = v
        t.settings = s
    ctx.audit("company.update", t.id)
    ctx.db.commit()
    return get_company(ctx)


# --------------------------------------------------------------------------- uploads & mapping review
@router.post("/library/{kind}/upload")
async def upload(kind: str, file: UploadFile = File(...), ctx: Ctx = Depends(require_roles("presales", "architect", "finance"))):
    if kind not in KINDS and kind != "template":
        raise HTTPException(404, "Unknown library section.")
    doc, data = await save_upload(ctx, file, kind)
    if kind == "template":
        if Path(file.filename).suffix.lower() != ".docx":
            raise HTTPException(400, "Proposal templates must be Word (.docx) files.")
        from docx import Document as Docx
        try:
            d = Docx(io.BytesIO(data))
        except Exception:
            raise HTTPException(400, "This Word file couldn't be opened. Save it again as .docx and retry.")
        text = "\n".join(p.text for p in d.paragraphs)
        hdr = "\n".join(p.text for s in d.sections for p in s.header.paragraphs)
        found = sorted(set(re.findall(r"\{\{([A-Z_]+)\}\}", text + hdr)))
        styles = [s.name for s in d.styles if s.type == 1 and s.name.startswith(("Heading", "Title", "Normal", "List"))][:12]
        # older templates become the history; newest wins
        for old in ctx.db.scalars(select(Document).where(Document.tenant_id == ctx.tenant_id, Document.kind == "template", Document.status == "imported")):
            old.status = "replaced"
        doc.status = "imported"
        from ..services.docgen import PLACEHOLDER_RE, template_mode
        sections = [p.text.strip() for p in d.paragraphs if PLACEHOLDER_RE.match(p.text or "")]
        mode = template_mode(d)
        how = {"marker": "The proposal is written where {{CONTENT}} sits.",
               "sections": f"Your template's own sections are kept: {len(sections)} placeholders are filled with the matching content, and its sample tables are left out.",
               "styles": "Only the look is used (styles, header, footer, logo); the proposal content replaces the body."}[mode]
        doc.parsed = {"placeholders": found, "styles": styles, "has_content_marker": "CONTENT" in found, "mode": mode,
                      "section_placeholders": sections[:40], "how": how}
        ctx.audit("template.upload", file.filename)
        ctx.db.commit()
        return {"id": doc.id, "kind": kind, "filename": doc.filename, "status": doc.status, "analysis": doc.parsed}
    try:
        parsed = parse_document(data, file.filename)
    except Exception as e:
        doc.status, doc.error = "failed", str(e)
        ctx.db.commit()
        raise HTTPException(400, f"Couldn't read this file: {e}")
    a = analyze(kind, parsed)
    stored = {k: v for k, v in a.items() if k != "records"}
    stored["records_preview"] = a["records"][:200]
    stored["record_count"] = len(a["records"])
    if parsed.tables:
        stored["rows"] = max(parsed.tables, key=len)[:2000]
    doc.parsed = stored
    doc.pages = parsed.pages
    doc.status = "review"
    ctx.db.commit()
    return {"id": doc.id, "kind": kind, "filename": doc.filename, "status": doc.status, "analysis": {**stored, "rows": None}}


@router.post("/library/uploads/{doc_id}/remap")
def remap(doc_id: str, body: dict, ctx: Ctx = Depends(require_roles("presales", "architect", "finance"))):
    doc = ctx.db.get(Document, doc_id)
    if not doc or doc.tenant_id != ctx.tenant_id:
        raise HTTPException(404, "Upload not found.")
    p = dict(doc.parsed or {})
    if p.get("format") != "table":
        raise HTTPException(400, "Column mapping only applies to spreadsheets.")
    mapping = {k: int(v) for k, v in (body.get("mapping") or {}).items() if v is not None and str(v) != ""}
    p["mapping"] = mapping
    if "wide_skill_columns" in body:
        p["wide_skill_columns"] = [int(x) for x in body["wide_skill_columns"]]
    recs = records_from_table(doc.kind, p["rows"], p)
    p["missing"] = missing_fields(doc.kind, p)
    p["records_preview"] = recs[:200]
    p["record_count"] = len(recs)
    doc.parsed = p
    ctx.db.commit()
    return {"analysis": {**p, "rows": None}}


def _import_doc(ctx: Ctx, doc: Document, mode: str = "replace", records: list | None = None) -> int:
    """Turn one reviewed upload into library records. The caller rebuilds the graph and commits."""
    kind = doc.kind
    p = doc.parsed or {}
    if p.get("format") == "table":
        recs = records_from_table(kind, p["rows"], p)
    else:
        recs = records or p.get("records_preview", [])
    if not recs:
        raise HTTPException(400, "Nothing to import. Check the column mapping.")
    M = MODEL[kind]
    if mode == "replace":
        ctx.db.execute(delete(M).where(M.tenant_id == ctx.tenant_id))
    cols = {c.name for c in M.__table__.columns}
    n = 0
    for r in recs:
        clean = {k: v for k, v in r.items() if k in cols and k not in ("id", "tenant_id")}
        if kind == "skills" and not clean.get("code"):
            clean["code"] = f"EMP-{n + 1:04d}"
        if kind == "skills" and (ctx.tenant.settings or {}).get("hide_employee_names", True):
            clean["name"] = ""
        if kind == "rates" and not clean.get("currency"):
            clean["currency"] = ctx.tenant.currency
        ctx.db.add(M(tenant_id=ctx.tenant_id, source_doc=doc.id, **clean))
        n += 1
    doc.status = "imported"
    ob = dict(ctx.tenant.onboarding or {})
    ob[kind] = True
    ctx.tenant.onboarding = ob
    return n


@router.post("/library/uploads/{doc_id}/import")
def import_upload(doc_id: str, body: dict | None = None, ctx: Ctx = Depends(require_roles("presales", "architect", "finance"))):
    body = body or {}
    doc = ctx.db.get(Document, doc_id)
    if not doc or doc.tenant_id != ctx.tenant_id:
        raise HTTPException(404, "Upload not found.")
    if doc.kind not in MODEL:
        raise HTTPException(400, "Pick what kind of file this is first.")
    n = _import_doc(ctx, doc, body.get("mode", "replace"), body.get("records"))
    ctx.db.flush()
    stats = graph.rebuild_company_graph(ctx.db, ctx.tenant_id)
    ctx.audit("library.import", doc.kind, records=n, file=doc.filename)
    ctx.db.commit()
    return {"imported": n, "kind": doc.kind, "graph": stats, "onboarding": _onboarding_status(ctx)}


# --------------------------------------------------------------------------- whole-folder import
SKIP_NAME = re.compile(r"^(~\$|\.|thumbs\.db$|desktop\.ini$)", re.I)
MAX_FOLDER_FILES = 300


def _batch_docs(ctx: Ctx, batch: str) -> list[Document]:
    rows = ctx.db.scalars(select(Document).where(Document.tenant_id == ctx.tenant_id, Document.kind != "rfp")
                          .order_by(Document.created_at)).all()
    return [d for d in rows if (d.parsed or {}).get("batch") == batch]


def _folder_row(d: Document) -> dict:
    p = d.parsed or {}
    return {"id": d.id, "path": p.get("path") or d.filename, "filename": d.filename, "size": d.size, "status": d.status,
            "type": d.kind if d.kind in folder_import.TYPES else None, "detected": p.get("detected"),
            "confidence": p.get("confidence"), "reason": p.get("reason", ""), "records": p.get("record_count", 0),
            "format": p.get("format"), "missing": p.get("missing", []), "error": d.error, "how": p.get("how"), "ai": p.get("ai")}


@router.post("/library/folder")
async def folder_upload(files: list[UploadFile] = File(...), paths: list[str] = Form(default=[]), batch: str = Form(default=""),
                        ctx: Ctx = Depends(require_roles("presales", "architect", "finance"))):
    """Receive some (or all) files of a folder. Large folders arrive in several calls with the same batch id."""
    batch = re.sub(r"[^a-z0-9]", "", batch.lower())[:32] or secrets.token_hex(8)
    existing = _batch_docs(ctx, batch)
    if len(existing) + len(files) > MAX_FOLDER_FILES:
        raise HTTPException(400, f"That's more than {MAX_FOLDER_FILES} files. Pick the folder that holds just the company knowledge files.")
    seen = {(d.filename, d.size) for d in existing}
    accepted, skipped = [], []
    for i, f in enumerate(files):
        rel = (paths[i] if i < len(paths) else "") or f.filename or "file"
        name = Path(rel).name
        ext = Path(name).suffix.lower()
        if SKIP_NAME.search(name):
            continue
        if ext not in ALLOWED:
            skipped.append({"path": rel, "reason": f"{ext or 'This type'} isn't a supported file"})
            continue
        data = await f.read()
        if not data:
            skipped.append({"path": rel, "reason": "Empty file"})
            continue
        if len(data) > settings.max_upload_mb * 1024 * 1024:
            skipped.append({"path": rel, "reason": f"Larger than {settings.max_upload_mb} MB"})
            continue
        if (name, len(data)) in seen:
            skipped.append({"path": rel, "reason": "Same file appears twice in the folder"})
            continue
        seen.add((name, len(data)))
        doc = Document(tenant_id=ctx.tenant_id, kind="unsorted", filename=name, path="", size=len(data), uploaded_by=ctx.user.id,
                       status="analyzing", parsed={"batch": batch, "path": rel})
        ctx.db.add(doc)
        ctx.db.flush()
        folder = UPLOAD_DIR / ctx.tenant_id
        folder.mkdir(parents=True, exist_ok=True)
        dest = folder / f"{doc.id}_{re.sub(r'[^A-Za-z0-9._ -]+', '_', name)[-120:]}"
        dest.write_bytes(data)
        doc.path = str(dest)
        accepted.append(doc.id)
    ctx.audit("library.folder_upload", batch, files=len(accepted))
    ctx.db.commit()
    folder_import.start(accepted)
    return {"batch": batch, "accepted": len(accepted), "skipped": skipped}


@router.get("/library/folder/{batch}")
def folder_status(batch: str, ctx: Ctx = Depends(current_ctx)):
    docs = _batch_docs(ctx, batch)
    rows = [_folder_row(d) for d in docs]
    return {"batch": batch, "files": rows, "pending": sum(1 for r in rows if r["status"] == "analyzing"),
            "imported": sum(1 for r in rows if r["status"] == "imported")}


@router.post("/library/folder/{batch}/{doc_id}/type")
def folder_set_type(batch: str, doc_id: str, body: dict, ctx: Ctx = Depends(require_roles("presales", "architect", "finance"))):
    doc = ctx.db.get(Document, doc_id)
    if not doc or doc.tenant_id != ctx.tenant_id or (doc.parsed or {}).get("batch") != batch:
        raise HTTPException(404, "File not found.")
    if doc.status in ("analyzing", "imported"):
        raise HTTPException(409, "This file is still being read." if doc.status == "analyzing" else "This file is already imported.")
    kind = body.get("type") or None
    if kind is not None and kind not in folder_import.TYPES:
        raise HTTPException(400, "Unknown file type.")
    try:
        folder_import.analyse_as(doc, kind, Path(doc.path).read_bytes(), None)
    except Exception as e:
        doc.status, doc.error = "failed", str(e)[:500]
    ctx.db.commit()
    if (doc.parsed or {}).get("ai") == "pending":
        folder_import.queue_ai(doc.id)
    return _folder_row(doc)


@router.post("/library/folder/{batch}/import")
def folder_import_all(batch: str, body: dict | None = None, ctx: Ctx = Depends(require_roles("presales", "architect", "finance"))):
    body = body or {}
    skip = set(body.get("skip") or [])
    replace = bool(body.get("replace", True))
    docs = [d for d in _batch_docs(ctx, batch) if d.status == "review"]
    if any(d.status == "analyzing" for d in _batch_docs(ctx, batch)):
        raise HTTPException(409, "Some files are still being read. Wait a moment and try again.")
    counts: dict[str, int] = {}
    files, problems = 0, []
    started: set[str] = set()
    template_done = False
    for d in docs:
        if d.id in skip or d.kind not in folder_import.TYPES:
            d.status = "skipped"
            continue
        if d.kind == "template":
            if template_done or d.error:
                d.status = "skipped"
                continue
            for old in ctx.db.scalars(select(Document).where(Document.tenant_id == ctx.tenant_id, Document.kind == "template",
                                                             Document.status == "imported")):
                old.status = "replaced"
            d.status, template_done = "imported", True
            counts["template"] = 1
            files += 1
            continue
        mode = "replace" if replace and d.kind not in started else "append"
        try:
            n = _import_doc(ctx, d, mode)
        except HTTPException as e:
            problems.append({"path": (d.parsed or {}).get("path") or d.filename, "reason": e.detail})
            continue
        started.add(d.kind)
        counts[d.kind] = counts.get(d.kind, 0) + n
        files += 1
    ctx.db.flush()
    stats = graph.rebuild_company_graph(ctx.db, ctx.tenant_id)
    ctx.audit("library.folder_import", batch, files=files, **counts)
    ctx.db.commit()
    return {"files": files, "counts": counts, "problems": problems, "graph": stats, "onboarding": _onboarding_status(ctx)}


@router.get("/library/uploads")
def list_uploads(ctx: Ctx = Depends(current_ctx)):
    rows = ctx.db.scalars(select(Document).where(Document.tenant_id == ctx.tenant_id, Document.kind.notin_(("rfp", "unsorted", "story_video")),
                                                 Document.status.notin_(("analyzing", "skipped"))).order_by(Document.created_at.desc())).all()
    return [{"id": d.id, "kind": d.kind, "filename": d.filename, "status": d.status, "size": d.size, "created_at": d.created_at,
             "records": (d.parsed or {}).get("record_count")} for d in rows]


@router.get("/library/uploads/{doc_id}")
def get_upload(doc_id: str, ctx: Ctx = Depends(current_ctx)):
    doc = ctx.db.get(Document, doc_id)
    if not doc or doc.tenant_id != ctx.tenant_id:
        raise HTTPException(404, "Upload not found.")
    return {"id": doc.id, "kind": doc.kind, "filename": doc.filename, "status": doc.status, "analysis": {**(doc.parsed or {}), "rows": None}}


@router.get("/library/template")
def get_template(ctx: Ctx = Depends(current_ctx)):
    doc = ctx.db.scalar(select(Document).where(Document.tenant_id == ctx.tenant_id, Document.kind == "template", Document.status == "imported")
                        .order_by(Document.created_at.desc()))
    if not doc:
        return None
    return {"id": doc.id, "filename": doc.filename, "created_at": doc.created_at, "analysis": doc.parsed}


@router.get("/library/template/download")
def download_template(ctx: Ctx = Depends(current_ctx)):
    doc = ctx.db.scalar(select(Document).where(Document.tenant_id == ctx.tenant_id, Document.kind == "template", Document.status == "imported")
                        .order_by(Document.created_at.desc()))
    if not doc:
        raise HTTPException(404, "No template uploaded yet.")
    return StreamingResponse(open(doc.path, "rb"), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             headers={"Content-Disposition": f'attachment; filename="{doc.filename}"'})


@router.get("/library/{kind}/starter.xlsx")
def starter(kind: str, ctx: Ctx = Depends(current_ctx)):
    """A blank, well-labelled spreadsheet for companies that want a clean start."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    if kind not in FIELDS:
        raise HTTPException(404, "Unknown library section.")
    headers = {
        "capabilities": ["Service Line", "Capability", "Description", "Technologies", "Industries", "Maturity"],
        "skills": ["Emp ID", "Designation", "Grade", "Total Exp (yrs)", "Base Location", "Practice", "Primary Skills (rating 1-5)", "Certifications", "Available %"],
        "certifications": ["Certification", "Type", "Issued By", "No. of Certified Employees", "Valid Until"],
        "rates": ["Role", "Level", "Delivery Location", "Day Rate", "Day Cost", "Currency"],
        "case_studies": ["Project Name", "Client", "Industry", "Region", "Year", "Summary", "Technologies", "Outcomes"],
        "accelerators": ["Accelerator", "Description", "Technologies", "Effort Saving %"],
    }[kind]
    example = {
        "capabilities": ["Integration Services", "API Management", "Design and run secure API platforms", "Apigee, MuleSoft", "BFSI; Healthcare", "Core"],
        "skills": ["E-001", "Senior Integration Engineer", "L4", 9, "Bengaluru", "Integration", "MuleSoft (5), Kafka (4), Java (4)", "MuleSoft Certified Developer", "100%"],
        "certifications": ["ISO/IEC 27001:2022", "Company", "BSI", "", "2027-06-30"],
        "rates": ["Integration Engineer", "L4", "Offshore", 380, 175, "USD"],
        "case_studies": ["API platform for a retail bank", "Retail bank", "BFSI", "UAE", 2025, "What we did and how", "MuleSoft, Kafka", "65 APIs live; 4x faster partner onboarding"],
        "accelerators": ["Interface Factory", "Reusable templates for common flows", "MuleSoft", 20],
    }[kind]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = KINDS[kind][:30]
    ws.append(headers)
    ws.append(example)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="125C7E")
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 26
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="revenueforge-{kind}-starter.xlsx"'})


# --------------------------------------------------------------------------- library CRUD
@router.get("/library/{kind}")
def list_kind(kind: str, q: str = "", ctx: Ctx = Depends(current_ctx)):
    if kind not in MODEL:
        raise HTTPException(404, "Unknown library section.")
    M = MODEL[kind]
    rows = ctx.db.scalars(select(M).where(M.tenant_id == ctx.tenant_id)).all()
    out = [_row_dict(kind, r) for r in rows]
    if q:
        ql = q.lower()
        out = [o for o in out if ql in str(o).lower()]
    hide = (ctx.tenant.settings or {}).get("hide_employee_names", True)
    if kind == "skills" and hide:
        for o in out:
            o["name"] = ""
    return out


@router.post("/library/{kind}")
def create_item(kind: str, body: dict, ctx: Ctx = Depends(require_roles("presales", "architect", "finance"))):
    if kind not in MODEL:
        raise HTTPException(404, "Unknown library section.")
    M = MODEL[kind]
    cols = {c.name for c in M.__table__.columns} - {"id", "tenant_id"}
    obj = M(tenant_id=ctx.tenant_id, **{k: v for k, v in body.items() if k in cols})
    ctx.db.add(obj)
    ctx.db.flush()
    graph.rebuild_company_graph(ctx.db, ctx.tenant_id)
    ctx.db.commit()
    return _row_dict(kind, obj)


@router.put("/library/{kind}/{item_id}")
def update_item(kind: str, item_id: str, body: dict, ctx: Ctx = Depends(require_roles("presales", "architect", "finance"))):
    if kind not in MODEL:
        raise HTTPException(404, "Unknown library section.")
    obj = ctx.db.get(MODEL[kind], item_id)
    if not obj or obj.tenant_id != ctx.tenant_id:
        raise HTTPException(404, "Item not found.")
    cols = {c.name for c in MODEL[kind].__table__.columns} - {"id", "tenant_id"}
    for k, v in body.items():
        if k in cols:
            setattr(obj, k, v)
    ctx.db.flush()
    graph.rebuild_company_graph(ctx.db, ctx.tenant_id)
    ctx.db.commit()
    return _row_dict(kind, obj)


@router.delete("/library/{kind}/{item_id}")
def delete_item(kind: str, item_id: str, ctx: Ctx = Depends(require_roles("presales", "architect", "finance"))):
    if kind not in MODEL:
        raise HTTPException(404, "Unknown library section.")
    obj = ctx.db.get(MODEL[kind], item_id)
    if not obj or obj.tenant_id != ctx.tenant_id:
        raise HTTPException(404, "Item not found.")
    ctx.db.delete(obj)
    ctx.db.flush()
    graph.rebuild_company_graph(ctx.db, ctx.tenant_id)
    ctx.db.commit()
    return {"ok": True}
