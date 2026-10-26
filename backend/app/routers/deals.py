"""Deals: upload an RFP, follow the pipeline, review and edit each stage, approve, download outputs."""
from __future__ import annotations

import io
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select

from ..models import STAGES, Deal, DealStage, Document, Requirement, RequirementComment, User
from ..security import Ctx, current_ctx, require_roles
from ..services import knowledge as K
from ..services import learning, pipeline
from ..services.masking import deal_masker, guess_client
from ..services.parsing import parse_document
from .library import save_upload

router = APIRouter(prefix="/api/deals", tags=["deals"])

STAGE_ROLES = {"requirements": ("presales",), "matching": ("presales", "architect"), "gaps": ("presales", "architect"),
               "architecture": ("architect",), "estimation": ("architect", "delivery"), "pricing": ("finance",),
               "proposal": ("presales",), "handoff": ("delivery", "presales")}


def _deal(ctx: Ctx, deal_id: str) -> Deal:
    d = ctx.db.get(Deal, deal_id)
    if not d or d.tenant_id != ctx.tenant_id:
        raise HTTPException(404, "Deal not found.")
    return d


def deal_out(ctx: Ctx, d: Deal, full: bool = False) -> dict:
    mask = deal_masker(d)
    stages = {s.stage: s for s in ctx.db.scalars(select(DealStage).where(DealStage.deal_id == d.id))}
    req_count = ctx.db.scalar(select(func.count()).select_from(Requirement).where(Requirement.deal_id == d.id))
    gaps = stages.get("gaps").data if "gaps" in stages else {}
    pricing = stages.get("pricing").data if "pricing" in stages else {}
    sel = None
    if pricing and pricing.get("scenarios"):
        sel = next((s for s in pricing["scenarios"] if s["name"] == pricing.get("selected", "Target")), None)
    out = {
        "id": d.id, "name": mask(d.name), "client": d.client_alias if d.confidential else d.client_name,
        "client_name": d.client_name, "client_alias": d.client_alias, "client_aliases": d.client_aliases, "confidential": d.confidential,
        "industry": d.industry, "sector": d.sector, "region": d.region, "doc_type": d.doc_type, "engagement_type": d.engagement_type,
        "currency": d.currency, "due_date": d.due_date, "status": d.status, "progress": d.progress or {}, "created_at": d.created_at,
        "updated_at": d.updated_at, "requirement_count": req_count, "final_value": d.final_value, "outcome_notes": d.outcome_notes,
        "coverage_pct": gaps.get("coverage_pct"), "readiness": gaps.get("readiness"), "mitigated_readiness": gaps.get("mitigated_readiness"),
        "verdict": gaps.get("verdict"),
        "value": round(sel["price"] * pricing.get("fx_to_rfp", 1), 2) if sel else None, "value_currency": pricing.get("rfp_currency") if sel else None,
        "margin_pct": sel["margin_pct"] if sel else None,
        "stages": [{"stage": s, "status": stages[s].status if s in stages else "pending", "approved_by": stages[s].approved_by if s in stages else "",
                    "approved_at": stages[s].approved_at if s in stages else None, "duration_ms": stages[s].duration_ms if s in stages else 0}
                   for s in STAGES],
    }
    if full:
        out["meta"] = mask(d.meta or {})
    return out


@router.get("")
def list_deals(ctx: Ctx = Depends(current_ctx)):
    rows = ctx.db.scalars(select(Deal).where(Deal.tenant_id == ctx.tenant_id).order_by(Deal.created_at.desc())).all()
    return [deal_out(ctx, d) for d in rows]


@router.post("/detect")
async def detect(file: UploadFile = File(...), ctx: Ctx = Depends(current_ctx)):
    """Quick look at an RFP before creating the deal: title, issuer and type, to prefill the form."""
    data = await file.read()
    try:
        parsed = parse_document(data, file.filename)
    except Exception as e:
        raise HTTPException(400, f"Couldn't read this file: {e}")
    from ..services.extraction import extract_metadata
    meta = extract_metadata(parsed)
    client, aliases = guess_client(parsed.text)
    return {"title": meta.get("title"), "doc_type": meta.get("doc_type"), "client_name": client, "client_aliases": aliases,
            "industry": meta.get("industry"), "region": meta.get("region"), "sector": meta.get("sector"), "pages": parsed.pages,
            "engagement_type": meta.get("engagement_type"), "currency": meta.get("currency"), "due_date": meta.get("due_date")}


@router.post("")
async def create_deal(file: UploadFile = File(...), name: str = Form(""), client_name: str = Form(""), client_alias: str = Form(""),
                      client_aliases: str = Form(""), confidential: bool = Form(True), engagement_type: str = Form(""),
                      ctx: Ctx = Depends(require_roles("presales"))):
    doc, data = await save_upload(ctx, file, "rfp")
    aliases = [a.strip() for a in client_aliases.split(",") if a.strip()]
    n = ctx.db.scalar(select(func.count()).select_from(Deal).where(Deal.tenant_id == ctx.tenant_id)) or 0
    d = Deal(tenant_id=ctx.tenant_id, name=name.strip() or "Untitled deal", client_name=client_name.strip(),
             client_alias=client_alias.strip() or f"Client {chr(65 + n % 26)}", client_aliases=aliases, confidential=confidential,
             engagement_type=engagement_type if engagement_type in ("managed_services", "implementation", "product", "consulting") else "",
             owner_id=ctx.user.id, rfp_document_id=doc.id, status="processing",
             progress={"stage": "requirements", "pct": 0, "message": "Queued"})
    ctx.db.add(d)
    ctx.audit("deal.create", d.name, file=file.filename)
    ctx.db.commit()
    pipeline.start(d.id, notify_user=ctx.user.id)
    return deal_out(ctx, d)


@router.get("/{deal_id}")
def get_deal(deal_id: str, ctx: Ctx = Depends(current_ctx)):
    return deal_out(ctx, _deal(ctx, deal_id), full=True)


@router.patch("/{deal_id}")
def update_deal(deal_id: str, body: dict, ctx: Ctx = Depends(require_roles("presales"))):
    d = _deal(ctx, deal_id)
    for f in ("name", "client_name", "client_alias", "industry", "region", "sector", "due_date", "currency", "outcome_notes"):
        if f in body:
            setattr(d, f, str(body[f] or ""))
    if "client_aliases" in body:
        d.client_aliases = [a.strip() for a in (body["client_aliases"] if isinstance(body["client_aliases"], list) else str(body["client_aliases"]).split(",")) if a.strip()]
    if "confidential" in body:
        d.confidential = bool(body["confidential"])
    if "final_value" in body:
        d.final_value = float(body["final_value"] or 0)
    if "actual_effort_days" in body:
        d.actual_effort_days = float(body["actual_effort_days"] or 0)
    rerun = None
    if "engagement_type" in body and body["engagement_type"] != d.engagement_type:
        if body["engagement_type"] not in ("managed_services", "implementation", "product", "consulting"):
            raise HTTPException(400, "Unknown engagement type.")
        d.engagement_type = body["engagement_type"]
        rerun = "matching"
    if "status" in body:
        if body["status"] not in ("in_review", "submitted", "won", "lost", "no_bid"):
            raise HTTPException(400, "Unknown status.")
        d.status = body["status"]
        ctx.audit("deal.outcome", d.id, status=d.status)
    ctx.db.commit()
    if rerun:
        pipeline.start(d.id, rerun)
    if body.get("status") in ("won", "lost") or "actual_effort_days" in body:
        from ..services import graph
        graph.add_deal_to_graph(ctx.db, d)
        learning.record_outcome(ctx.db, d)
        learning.train_classifier(ctx.db, ctx.tenant_id)
        ctx.db.commit()
    return deal_out(ctx, d, full=True)


@router.delete("/{deal_id}")
def delete_deal(deal_id: str, ctx: Ctx = Depends(require_roles("presales"))):
    d = _deal(ctx, deal_id)
    from sqlalchemy import delete as sqldelete
    from ..models import GraphEdge, GraphNode
    keys = [n.key for n in ctx.db.scalars(select(GraphNode).where(GraphNode.tenant_id == ctx.tenant_id)) if (n.props or {}).get("deal_id") == d.id]
    if keys:
        ctx.db.execute(sqldelete(GraphEdge).where(GraphEdge.tenant_id == ctx.tenant_id, (GraphEdge.src.in_(keys)) | (GraphEdge.dst.in_(keys))))
        ctx.db.execute(sqldelete(GraphNode).where(GraphNode.tenant_id == ctx.tenant_id, GraphNode.key.in_(keys)))
    ctx.db.delete(d)
    ctx.audit("deal.delete", deal_id)
    ctx.db.commit()
    return {"ok": True}


@router.post("/{deal_id}/rerun")
def rerun(deal_id: str, body: dict | None = None, ctx: Ctx = Depends(require_roles("presales", "architect", "finance", "delivery"))):
    d = _deal(ctx, deal_id)
    stage = (body or {}).get("from_stage", "requirements")
    if stage not in STAGES:
        raise HTTPException(400, "Unknown stage.")
    d.status = "processing"
    d.progress = {"stage": stage, "pct": int(STAGES.index(stage) / len(STAGES) * 100), "message": "Queued"}
    ctx.db.commit()
    pipeline.start(d.id, stage, notify_user=ctx.user.id)
    return deal_out(ctx, d)


# --------------------------------------------------------------------------- requirements
def _req_out(r: Requirement, mask) -> dict:
    return {"id": r.id, "code": r.code, "source_ref": r.source_ref, "text": mask(r.text), "category": r.category, "subcategory": r.subcategory,
            "priority": r.priority, "page": r.page, "section": mask(r.section), "block_idx": r.block_idx, "class_confidence": r.class_confidence,
            "technologies": r.technologies or [], "match_status": r.match_status, "match_score": r.match_score, "match": mask(r.match or {}),
            "gap": r.gap or {}, "workstream": r.workstream, "complexity": r.complexity, "effort_days": r.effort_days, "verified": r.verified,
            "edited": r.edited, "original_category": r.original_category, "notes": r.notes, "classified_by": r.classified_by,
            "response": mask(r.response or ""), "response_source": r.response_source}


@router.get("/{deal_id}/requirements")
def list_requirements(deal_id: str, ctx: Ctx = Depends(current_ctx)):
    d = _deal(ctx, deal_id)
    mask = deal_masker(d)
    rows = ctx.db.scalars(select(Requirement).where(Requirement.deal_id == d.id).order_by(Requirement.seq)).all()
    counts = dict(ctx.db.execute(select(RequirementComment.requirement_id, func.count()).where(RequirementComment.deal_id == d.id)
                                 .group_by(RequirementComment.requirement_id)).all())
    return [{**_req_out(r, mask), "comments": counts.get(r.id, 0)} for r in rows]


# --------------------------------------------------------------------------- comments on requirements
def _comment_out(c: RequirementComment, mask, me: str) -> dict:
    return {"id": c.id, "author": c.author, "user_id": c.user_id, "text": mask(c.text), "created_at": c.created_at, "mine": c.user_id == me}


def _mentions(ctx: Ctx, text: str) -> list[str]:
    """@name → user ids: matches a first name, a full name written without spaces, or the part of an email before @."""
    tags = {t.lower() for t in re.findall(r"@([A-Za-z0-9._-]{2,60})", text)}
    if not tags:
        return []
    ids = []
    for u in ctx.db.scalars(select(User).where(User.tenant_id == ctx.tenant_id, User.is_active == True)):  # noqa: E712
        keys = {u.email.split("@")[0].lower(), (u.name or "").split(" ")[0].lower(), re.sub(r"\s+", "", (u.name or "")).lower(),
                re.sub(r"\s+", ".", (u.name or "").strip()).lower()}
        if tags & keys:
            ids.append(u.id)
    return ids


@router.get("/{deal_id}/requirements/{req_id}/comments")
def list_comments(deal_id: str, req_id: str, ctx: Ctx = Depends(current_ctx)):
    d = _deal(ctx, deal_id)
    mask = deal_masker(d)
    rows = ctx.db.scalars(select(RequirementComment).where(RequirementComment.deal_id == d.id, RequirementComment.requirement_id == req_id)
                          .order_by(RequirementComment.created_at)).all()
    return [_comment_out(c, mask, ctx.user.id) for c in rows]


@router.post("/{deal_id}/requirements/{req_id}/comments")
def add_comment(deal_id: str, req_id: str, body: dict, ctx: Ctx = Depends(current_ctx)):
    d = _deal(ctx, deal_id)
    r = ctx.db.get(Requirement, req_id)
    if not r or r.deal_id != d.id:
        raise HTTPException(404, "Requirement not found.")
    text = str(body.get("text") or "").strip()
    if not text:
        raise HTTPException(400, "Write a comment first.")
    ids = _mentions(ctx, text[:4000])
    c = RequirementComment(tenant_id=ctx.tenant_id, deal_id=d.id, requirement_id=r.id, user_id=ctx.user.id, author=ctx.user.name,
                           text=text[:4000], mentions=ids)
    ctx.db.add(c)
    ctx.audit("requirement.comment", f"{d.id}:{r.code}", mentions=len(ids))
    ctx.db.commit()
    if ids:
        from ..services import notify
        notify.mentioned(ctx.db, d, ctx.user, c.text, r.code, ids)
    return _comment_out(c, deal_masker(d), ctx.user.id)


@router.delete("/{deal_id}/comments/{comment_id}")
def delete_comment(deal_id: str, comment_id: str, ctx: Ctx = Depends(current_ctx)):
    d = _deal(ctx, deal_id)
    c = ctx.db.get(RequirementComment, comment_id)
    if not c or c.deal_id != d.id:
        raise HTTPException(404, "Comment not found.")
    if c.user_id != ctx.user.id and ctx.user.role != "admin":
        raise HTTPException(403, "You can only delete your own comments.")
    ctx.db.delete(c)
    ctx.db.commit()
    return {"ok": True}


@router.patch("/{deal_id}/requirements/{req_id}")
def update_requirement(deal_id: str, req_id: str, body: dict, ctx: Ctx = Depends(require_roles("presales", "architect"))):
    d = _deal(ctx, deal_id)
    r = ctx.db.get(Requirement, req_id)
    if not r or r.deal_id != d.id:
        raise HTTPException(404, "Requirement not found.")
    learn = lambda kind, ai, human: learning.record(ctx.db, ctx.tenant_id, kind, r.text, ai, human, deal_id=d.id, user=ctx.user.name,
                                                    context={"section": r.section, "code": r.code})
    if "category" in body:
        if body["category"] not in K.CATEGORIES:
            raise HTTPException(400, f"Category must be one of {', '.join(K.CATEGORIES)}.")
        if body["category"] != r.category:
            learn("category", r.category, body["category"])
            r.category = body["category"]
            if "subcategory" not in body:
                r.subcategory = K.best_subcategory(r.text, r.section, r.category)
            r.classified_by = "human"
    if "subcategory" in body and body["subcategory"] != r.subcategory:
        learn("subcategory", r.subcategory, body["subcategory"])
        r.subcategory = body["subcategory"]
    if "priority" in body and body["priority"] != r.priority:
        learn("priority", r.priority, body["priority"])
        r.priority = body["priority"]
    for f in ("text", "notes"):
        if f in body:
            setattr(r, f, body[f])
    if "match_status" in body:
        if body["match_status"] not in ("strong", "partial", "gap", "confirm"):
            raise HTTPException(400, "Unknown match status.")
        if body["match_status"] != r.match_status:
            learn("match", r.match_status, body["match_status"])
        r.match_status = body["match_status"]
        m = dict(r.match or {})
        m["overridden_by"] = ctx.user.name
        r.match = m
    if "response" in body:
        new = str(body["response"] or "").strip()
        if new != (r.response or "").strip():
            learn("response", r.response or "", new)
        r.response, r.response_source = new, "human"
    if "verified" in body:
        r.verified = bool(body["verified"])
    if any(f in body for f in ("category", "subcategory", "priority", "text", "match_status")):
        r.edited = True
        r.verified = True
    ctx.audit("requirement.update", r.code, fields=list(body.keys()))
    ctx.db.commit()
    return _req_out(r, deal_masker(d))


@router.post("/{deal_id}/requirements")
def add_requirement(deal_id: str, body: dict, ctx: Ctx = Depends(require_roles("presales"))):
    d = _deal(ctx, deal_id)
    text = (body.get("text") or "").strip()
    if len(text) < 5:
        raise HTTPException(400, "Write the requirement text.")
    seq = (ctx.db.scalar(select(func.max(Requirement.seq)).where(Requirement.deal_id == d.id)) or 0) + 1
    cat, sub, conf = K.classify(text, body.get("section", ""))
    r = Requirement(tenant_id=ctx.tenant_id, deal_id=d.id, seq=seq, code=f"REQ-{seq:03d}", text=text, category=body.get("category") or cat,
                    subcategory=body.get("subcategory") or sub, priority=body.get("priority") or K.priority(text), page=int(body.get("page") or 1),
                    section=body.get("section", "Added manually"), class_confidence=1.0, technologies=K.find_technologies(text),
                    original_category=body.get("category") or cat, verified=True, edited=True, classified_by="human")
    ctx.db.add(r)
    learning.record(ctx.db, ctx.tenant_id, "added", text, "", r.category, deal_id=d.id, user=ctx.user.name)
    ctx.db.commit()
    return _req_out(r, deal_masker(d))


@router.delete("/{deal_id}/requirements/{req_id}")
def delete_requirement(deal_id: str, req_id: str, ctx: Ctx = Depends(require_roles("presales"))):
    d = _deal(ctx, deal_id)
    r = ctx.db.get(Requirement, req_id)
    if not r or r.deal_id != d.id:
        raise HTTPException(404, "Requirement not found.")
    learning.record(ctx.db, ctx.tenant_id, "removed", r.text, r.category, "not a requirement", deal_id=d.id, user=ctx.user.name)
    ctx.db.delete(r)
    ctx.db.commit()
    return {"ok": True}


@router.get("/{deal_id}/accuracy")
def accuracy(deal_id: str, ctx: Ctx = Depends(current_ctx)):
    """Classification accuracy measured against human review: reviewed requirements whose category was kept."""
    d = _deal(ctx, deal_id)
    rows = ctx.db.scalars(select(Requirement).where(Requirement.deal_id == d.id)).all()
    reviewed = [r for r in rows if r.verified]
    correct = [r for r in reviewed if r.category == r.original_category]
    conf = Counter()
    for r in reviewed:
        conf[f"{r.original_category}→{r.category}"] += 1
    return {"total": len(rows), "reviewed": len(reviewed), "correct": len(correct),
            "accuracy": round(len(correct) / len(reviewed) * 100, 1) if reviewed else None,
            "confusion": dict(conf), "model_confidence": round(sum(r.class_confidence for r in rows) / max(1, len(rows)) * 100, 1)}


@router.post("/{deal_id}/accuracy/labels")
async def accuracy_labels(deal_id: str, file: UploadFile = File(...), ctx: Ctx = Depends(require_roles("presales"))):
    """Score against a labelled test set (xlsx/csv with columns: requirement text, category)."""
    d = _deal(ctx, deal_id)
    parsed = parse_document(await file.read(), file.filename)
    rows = max(parsed.tables, key=len) if parsed.tables else []
    if len(rows) < 2:
        raise HTTPException(400, "Upload a spreadsheet with a requirement column and a category column.")
    header = [h.lower() for h in rows[0]]
    ti = next((i for i, h in enumerate(header) if "requirement" in h or "text" in h), 0)
    ci = next((i for i, h in enumerate(header) if "categ" in h or "class" in h or "type" in h), 1)
    labels = [(r[ti], r[ci]) for r in rows[1:] if len(r) > max(ti, ci) and r[ti] and r[ci]]
    reqs = ctx.db.scalars(select(Requirement).where(Requirement.deal_id == d.id)).all()
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    vec = TfidfVectorizer().fit([r.text for r in reqs] + [l[0] for l in labels])
    sims = cosine_similarity(vec.transform([l[0] for l in labels]), vec.transform([r.text for r in reqs]))
    found = correct = 0
    misses = []
    for i, (txt, cat) in enumerate(labels):
        j = sims[i].argmax()
        if sims[i][j] >= 0.55:
            found += 1
            norm = cat.strip().lower().replace("-", "").replace(" ", "")
            if reqs[j].category.lower().replace("-", "") == norm:
                correct += 1
            else:
                misses.append({"label": cat, "predicted": reqs[j].category, "text": txt[:160]})
        else:
            misses.append({"label": cat, "predicted": "(not extracted)", "text": txt[:160]})
    return {"labelled": len(labels), "recall": round(found / len(labels) * 100, 1) if labels else 0,
            "classification_accuracy": round(correct / found * 100, 1) if found else 0, "misses": misses[:50]}


# --------------------------------------------------------------------------- source document
@router.get("/{deal_id}/source")
def source(deal_id: str, ctx: Ctx = Depends(current_ctx)):
    d = _deal(ctx, deal_id)
    doc = ctx.db.get(Document, d.rfp_document_id)
    if not doc:
        raise HTTPException(404, "Source document missing.")
    parsed = pipeline.get_parsed(ctx.db, doc)
    ctx.db.commit()
    mask = deal_masker(d)
    return {"filename": mask(doc.filename), "pages": parsed.pages,
            "blocks": [{"idx": b.idx, "page": b.page, "kind": b.kind, "text": mask(b.text), "ref": b.ref} for b in parsed.blocks]}


# --------------------------------------------------------------------------- stages
@router.get("/{deal_id}/stages/{stage}")
def get_stage(deal_id: str, stage: str, ctx: Ctx = Depends(current_ctx)):
    d = _deal(ctx, deal_id)
    if stage not in STAGES:
        raise HTTPException(404, "Unknown stage.")
    st = ctx.db.scalar(select(DealStage).where(DealStage.deal_id == d.id, DealStage.stage == stage))
    if not st:
        return {"stage": stage, "status": "pending", "data": {}}
    data = dict(st.data or {})
    data.pop("file", None)
    return {"stage": stage, "status": st.status, "data": deal_masker(d)(data), "approved_by": st.approved_by, "approved_at": st.approved_at,
            "duration_ms": st.duration_ms}


@router.post("/{deal_id}/stages/{stage}/approve")
def approve(deal_id: str, stage: str, ctx: Ctx = Depends(current_ctx)):
    d = _deal(ctx, deal_id)
    if stage not in STAGES:
        raise HTTPException(404, "Unknown stage.")
    allowed = STAGE_ROLES[stage]
    if ctx.user.role != "admin" and ctx.user.role not in allowed:
        raise HTTPException(403, f"Approving {stage} needs the {' or '.join(allowed)} role.")
    st = ctx.db.scalar(select(DealStage).where(DealStage.deal_id == d.id, DealStage.stage == stage))
    if not st or st.status not in ("draft", "approved"):
        raise HTTPException(400, "This stage has no draft to approve yet.")
    st.status = "approved"
    st.approved_by = ctx.user.name
    st.approved_at = datetime.now(timezone.utc)
    ctx.audit("stage.approve", f"{d.id}:{stage}")
    learned = {}
    if stage == "requirements":
        # Approval confirms the categories: they become training data for the company's own classifier.
        for r in ctx.db.scalars(select(Requirement).where(Requirement.deal_id == d.id)):
            r.verified = True
        ctx.db.flush()
        learned["classifier"] = learning.train_classifier(ctx.db, ctx.tenant_id)
    if stage == "proposal":
        learned["answers_added"] = learning.harvest_answers(ctx.db, d, ctx.user.name)
    ctx.db.commit()
    from ..services import notify
    notify.next_stage_ready(ctx.db, d, exclude_user_id=ctx.user.id, after=stage)  # "your turn" email to the next approvers
    return {"stage": stage, "status": st.status, "approved_by": st.approved_by, "approved_at": st.approved_at, "learned": learned}


@router.put("/{deal_id}/stages/{stage}/overrides")
def overrides(deal_id: str, stage: str, body: dict, ctx: Ctx = Depends(require_roles("architect", "finance", "delivery"))):
    """Save user adjustments (rates, margin, contingency, duration, selected scenario) and recompute downstream."""
    d = _deal(ctx, deal_id)
    if stage not in ("estimation", "pricing"):
        raise HTTPException(400, "Only estimation and pricing accept overrides.")
    st = ctx.db.scalar(select(DealStage).where(DealStage.deal_id == d.id, DealStage.stage == stage))
    if not st:
        raise HTTPException(400, "Run the pipeline first.")
    data = dict(st.data or {})
    ov = dict(data.get("overrides", {}))
    for k, v in (body.get("overrides") or {}).items():
        if v is None or v == "":
            ov.pop(k, None)
        else:
            ov[k] = v
    data["overrides"] = ov
    if "selected" in body:
        data["selected"] = body["selected"]
    st.data = data
    ctx.audit("stage.override", f"{d.id}:{stage}", overrides=ov)
    ctx.db.commit()
    # recompute this stage and everything after it synchronously (fast, no LLM involved until proposal)
    db = ctx.db
    for s in STAGES[STAGES.index(stage):STAGES.index("pricing") + 1]:
        pipeline.run_stage(db, d, s)
    d.status = "processing"
    db.commit()
    pipeline.start(d.id, "proposal")
    return get_stage(deal_id, stage, ctx)


@router.post("/{deal_id}/proposal/rebuild")
def rebuild_proposal(deal_id: str, ctx: Ctx = Depends(require_roles("presales"))):
    """Rebuild the Word documents from the current (edited) responses without re-running the AI."""
    d = _deal(ctx, deal_id)
    from ..services import docgen
    ctxd = pipeline.context(ctx.db, d)
    st0 = ctx.db.scalar(select(DealStage).where(DealStage.deal_id == d.id, DealStage.stage == "proposal"))
    if st0 and st0.data.get("executive_summary"):
        ctxd["executive_summary"] = st0.data["executive_summary"]
    data, info = docgen.build_proposal(ctxd, pipeline._template_bytes(ctx.db, ctx.tenant_id))
    st = ctx.db.scalar(select(DealStage).where(DealStage.deal_id == d.id, DealStage.stage == "proposal"))
    path = Path(st.data["file"]) if st and st.data.get("file") else None
    if path:
        path.write_bytes(data)
        st.data = {**st.data, **{k: v for k, v in info.items() if k != "executive_summary"}}
    ctx.db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------- proposal video
@router.get("/{deal_id}/story")
def story_status(deal_id: str, ctx: Ctx = Depends(current_ctx)):
    from ..services import story
    d = _deal(ctx, deal_id)
    st = ctx.db.scalar(select(DealStage).where(DealStage.deal_id == d.id, DealStage.stage == "estimation"))
    out = story.status(ctx.db, d)
    out["can_make"] = bool(st and st.data)
    out["has_client_name"] = bool((d.client_name or "").strip())
    return out


@router.post("/{deal_id}/story")
def story_make(deal_id: str, body: dict | None = None, ctx: Ctx = Depends(require_roles("presales"))):
    from ..services import story
    d = _deal(ctx, deal_id)
    if d.status == "processing":
        raise HTTPException(409, "Wait until the deal has finished processing.")
    st = ctx.db.scalar(select(DealStage).where(DealStage.deal_id == d.id, DealStage.stage == "estimation"))
    if not st or not st.data:
        raise HTTPException(400, "Run the deal up to Estimation first: the video shows the solution, team and plan.")
    if not story.ffmpeg_exe():
        raise HTTPException(503, "The video encoder isn't installed yet. Stop the app, run setup.bat, then start it again.")
    try:
        story.start(ctx.db, d, ctx.user.id, bool((body or {}).get("use_client_name", True)))
    except RuntimeError as e:
        raise HTTPException(409, str(e))
    ctx.audit("story.generate", d.id)
    ctx.db.commit()
    return story.status(ctx.db, d)


def _story_file(ctx: Ctx, d: Deal) -> Path:
    from ..services import story
    doc = story.find(ctx.db, ctx.tenant_id, d.id)
    if not doc or doc.status != "ready" or not Path(doc.path or "").exists():
        raise HTTPException(404, "The video hasn't been made yet.")
    return Path(doc.path)


# --------------------------------------------------------------------------- downloads
@router.get("/{deal_id}/download/{what}")
def download(deal_id: str, what: str, ctx: Ctx = Depends(current_ctx)):
    d = _deal(ctx, deal_id)
    safe = "".join(c for c in deal_masker(d)(d.name) if c.isalnum() or c in " -_")[:60].strip() or "deal"
    if what in ("proposal", "handoff"):
        st = ctx.db.scalar(select(DealStage).where(DealStage.deal_id == d.id, DealStage.stage == what))
        if not st or not (st.data or {}).get("file") or not Path(st.data["file"]).exists():
            raise HTTPException(404, "This document hasn't been generated yet.")
        ctx.audit("download", f"{d.id}:{what}")
        ctx.db.commit()
        return StreamingResponse(open(st.data["file"], "rb"), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                 headers={"Content-Disposition": f'attachment; filename="{safe} - {what}.docx"'})
    if what == "story":
        f = _story_file(ctx, d)
        ctx.audit("download", f"{d.id}:story")
        ctx.db.commit()
        return StreamingResponse(open(f, "rb"), media_type="video/mp4",
                                 headers={"Content-Disposition": f'attachment; filename="{safe} - solution video.mp4"', "Content-Length": str(f.stat().st_size)})
    if what == "proposal_pack":
        import zipfile
        st = ctx.db.scalar(select(DealStage).where(DealStage.deal_id == d.id, DealStage.stage == "proposal"))
        if not st or not (st.data or {}).get("file") or not Path(st.data["file"]).exists():
            raise HTTPException(404, "The proposal hasn't been generated yet.")
        f = _story_file(ctx, d)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(st.data["file"], f"{safe} - proposal.docx")
            z.write(f, f"{safe} - solution video.mp4", compress_type=zipfile.ZIP_STORED)
        buf.seek(0)
        ctx.audit("download", f"{d.id}:proposal_pack")
        ctx.db.commit()
        return StreamingResponse(buf, media_type="application/zip",
                                 headers={"Content-Disposition": f'attachment; filename="{safe} - proposal with video.zip"'})
    if what == "requirements":
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        mask = deal_masker(d)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Requirements"
        hdr = ["Code", "RFP ref", "Requirement", "Category", "Subcategory", "Priority", "Page", "Section", "Match", "Score", "Evidence",
               "Gap risk", "Mitigation", "Workstream", "Complexity", "Effort (days)", "Verified"]
        ws.append(hdr)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="125C7E")
        for r in ctx.db.scalars(select(Requirement).where(Requirement.deal_id == d.id).order_by(Requirement.seq)):
            ev = "; ".join(f"{e['type']}: {e['name']}" for e in (r.match or {}).get("evidence", [])[:3])
            opt = (r.gap or {}).get("options", [])
            ws.append([r.code, r.source_ref, mask(r.text), r.category, r.subcategory, r.priority, r.page, mask(r.section), r.match_status,
                       r.match_score, ev, (r.gap or {}).get("risk", ""), (opt[0]["type"] + ": " + opt[0]["detail"]) if opt else "",
                       r.workstream, r.complexity, r.effort_days, "Yes" if r.verified else ""])
        for i, w in enumerate([10, 10, 80, 14, 22, 11, 6, 30, 10, 8, 50, 9, 60, 24, 10, 10, 9], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f'attachment; filename="{safe} - requirements.xlsx"'})
    raise HTTPException(404, "Unknown download.")


# --------------------------------------------------------------------------- client clarification questions
@router.get("/{deal_id}/clarifications")
def get_clarifications(deal_id: str, ctx: Ctx = Depends(current_ctx)):
    d = _deal(ctx, deal_id)
    return (d.meta or {}).get("clarifications") or {"items": [], "by": None, "generated_at": None}


@router.post("/{deal_id}/clarifications")
def make_clarifications(deal_id: str, ctx: Ctx = Depends(require_roles("presales", "architect"))):
    from ..services import clarify
    d = _deal(ctx, deal_id)
    if d.status == "processing":
        raise HTTPException(409, "Wait until the deal has finished processing.")
    ctx.audit("deal.clarifications", d.id)
    return clarify.generate(ctx.db, d)


@router.put("/{deal_id}/clarifications")
def save_clarifications(deal_id: str, body: dict, ctx: Ctx = Depends(require_roles("presales", "architect"))):
    d = _deal(ctx, deal_id)
    items = [{"code": str(i.get("code") or "")[:40], "question": str(i.get("question") or "").strip()[:800], "why": str(i.get("why") or "")[:200]}
             for i in (body.get("items") or [])[:60] if str(i.get("question") or "").strip()]
    meta = dict(d.meta or {})
    cur = dict(meta.get("clarifications") or {})
    cur.update({"items": items, "edited_by": ctx.user.name, "edited_at": datetime.now(timezone.utc).isoformat()})
    meta["clarifications"] = cur
    d.meta = meta
    ctx.db.commit()
    return cur


# --------------------------------------------------------------------------- similar past deals
@router.get("/{deal_id}/similar")
def similar_deals(deal_id: str, ctx: Ctx = Depends(current_ctx)):
    from ..services import similar
    return similar.similar(ctx.db, _deal(ctx, deal_id))
